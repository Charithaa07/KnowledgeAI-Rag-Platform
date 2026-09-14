import io
import csv

from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation

SUPPORTED = {"pdf", "docx", "txt", "md", "markdown", "csv", "pptx", "ppt"}


def _from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "") for p in reader.pages)


def _from_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _from_pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _from_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="ignore")
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(r) for r in rows)


def _from_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


_HANDLERS = {
    "pdf": _from_pdf,
    "docx": _from_docx,
    "pptx": _from_pptx,
    "ppt": _from_pptx,
    "csv": _from_csv,
}


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    handler = _HANDLERS.get(ext, _from_plain)  # txt, md, markdown fall back to plain text
    return handler(data)


def chunk_text(text: str, size: int = 1000, overlap: int = 150):
    text = " ".join(text.split())
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        chunks.append(text[start:end])
        if end == n:
            break
        start = end - overlap
    return chunks
